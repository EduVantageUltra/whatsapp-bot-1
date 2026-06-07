import anthropic
import gspread
import base64
import os
import sys
import json
from google.oauth2.service_account import Credentials
from datetime import datetime
from config import ANTHROPIC_API_KEY, GOOGLE_CREDENTIALS_FILE, SHEET_NAME, SHEET_ID

HEADERS = [
    "Date", "Time", "Name", "Phone", "Email",
    "Company", "Address", "Invoice No",
    "Amount", "Extra Info", "Status",
    "Source File", "File Type"
]

def connect_to_sheet():
    print("Google Sheet se connect ho raha hoon...")
    scopes = [
        "https://www.googleapis.com/auth/spreadsheets",
        "https://www.googleapis.com/auth/drive"
    ]
    creds = Credentials.from_service_account_file(GOOGLE_CREDENTIALS_FILE, scopes=scopes)
    client = gspread.authorize(creds)
    sheet = client.open_by_key(SHEET_ID).sheet1

    first_row = sheet.row_values(1)
    if not first_row or first_row[0] != "Date":
        sheet.clear()
        sheet.append_row(HEADERS)
        sheet.format("A1:M1", {
            "textFormat": {"bold": True},
            "backgroundColor": {"red": 0.2, "green": 0.6, "blue": 1.0}
        })
        print("✅ Headers set ho gaye!")
    return sheet

def get_file_type(file_path):
    ext = os.path.splitext(file_path)[1].lower()
    if ext in [".jpg", ".jpeg", ".png", ".gif", ".webp", ".bmp", ".tiff"]:
        return "image"
    elif ext == ".pdf":
        return "pdf"
    elif ext in [".doc", ".docx"]:
        return "word"
    elif ext in [".xls", ".xlsx", ".csv"]:
        return "excel"
    return "unknown"

def get_extraction_prompt():
    return """You are an expert data extraction agent. Extract EVERY person, business entity, and contact from this document. Missing any entry is NOT acceptable — exhaustive extraction is the only goal.

DOCUMENT TYPES YOU WILL ENCOUNTER:
- Invoice/Bill/Receipt: Extract BOTH buyer AND seller as separate entries. Add any other parties too (transporter, broker, agent, consignee, etc.)
- Visiting card (single or multiple): One entry per card, all details.
- Handwritten notes/forms/registers: Read carefully even if writing is messy, blurry, or in mixed scripts. Make your best attempt; note uncertainty in extra_info.
- Table/list data: Each row that contains contact info = one separate entry in the array.
- Screenshots (WhatsApp, email, website, app): Extract every distinct person or business visible.
- Hindi/English/Hinglish text: Handle all three equally — Devanagari names, transliterated names, mixed text.

EXTRACTION RULES:
1. ALWAYS return a JSON array — even if only one entry found.
2. phone: digits and + only. Remove all spaces, dashes, brackets, dots. If multiple phone numbers exist for one entity, put the primary one in "phone" and rest in "extra_info".
3. amount: include currency symbol if visible (e.g., "Rs.2500", "₹1,23,456.00", "2500"). For invoices use the final payable/total amount.
4. gst_number: 15-character GSTIN (e.g., "07AABCU9603R1ZP"). Extract exactly as written.
5. date: convert to DD-MM-YYYY format. Use the document/invoice date, not today.
6. invoice_number: bill no., order no., receipt no., voucher no. — any document identifier.
7. company: full legal name preferred (Pvt Ltd, LLP, OPC, etc.). Trade name also fine.
8. address: include pincode, city, state if visible. Multi-line address as single string.
9. If a field is genuinely not present, use null. Do NOT guess names or phone numbers.
10. extra_info: use this for — role label (Buyer / Seller / Consignee / Agent / Transporter), alternate phone numbers, website, bank details, extra context, or uncertainty notes for handwritten text.
11. For table data with many rows: create one entry per row; share invoice_number/date across entries from the same document.
12. Blank or empty rows in tables: skip them.

OUTPUT — return ONLY the JSON array below, absolutely nothing else before or after:
[
    {
        "name": "full person name or contact person name or null",
        "phone": "digits and + sign only or null",
        "email": "email address or null",
        "company": "company or business name or null",
        "address": "complete address or null",
        "invoice_number": "invoice/bill/order/receipt number or null",
        "amount": "total payable amount with currency symbol or null",
        "gst_number": "GSTIN or null",
        "date": "DD-MM-YYYY or null",
        "extra_info": "role, alternate contacts, other relevant info or null"
    }
]

CRITICAL: Scan the ENTIRE document — top to bottom, including headers, footers, stamps, watermarks, side text, and any small-print areas. Every missed entry is a failure."""

def process_image_file(file_path):
    print("Image process ho rahi hai...")
    ext = os.path.splitext(file_path)[1].lower()
    format_map = {
        ".jpg": "image/jpeg", ".jpeg": "image/jpeg",
        ".png": "image/png", ".gif": "image/gif",
        ".webp": "image/webp", ".bmp": "image/png", ".tiff": "image/png"
    }
    media_type = format_map.get(ext, "image/jpeg")
    with open(file_path, "rb") as f:
        image_data = base64.standard_b64encode(f.read()).decode("utf-8")
    client = anthropic.Anthropic(api_key=ANTHROPIC_API_KEY)
    message = client.messages.create(
        model="claude-opus-4-5",
        max_tokens=4096,
        messages=[{
            "role": "user",
            "content": [
                {"type": "image", "source": {"type": "base64", "media_type": media_type, "data": image_data}},
                {"type": "text", "text": get_extraction_prompt()}
            ]
        }]
    )
    return message.content[0].text.strip()

def process_pdf_file(file_path):
    print("PDF process ho rahi hai...")
    with open(file_path, "rb") as f:
        pdf_data = base64.standard_b64encode(f.read()).decode("utf-8")
    client = anthropic.Anthropic(api_key=ANTHROPIC_API_KEY)
    message = client.messages.create(
        model="claude-opus-4-5",
        max_tokens=4096,
        messages=[{
            "role": "user",
            "content": [
                {"type": "document", "source": {"type": "base64", "media_type": "application/pdf", "data": pdf_data}},
                {"type": "text", "text": get_extraction_prompt()}
            ]
        }]
    )
    return message.content[0].text.strip()

def process_word_file(file_path):
    print("Word file process ho rahi hai...")
    try:
        from docx import Document
    except ImportError:
        os.system("pip install python-docx")
        from docx import Document
    doc = Document(file_path)
    text = "\n".join([p.text for p in doc.paragraphs if p.text.strip()])
    client = anthropic.Anthropic(api_key=ANTHROPIC_API_KEY)
    message = client.messages.create(
        model="claude-opus-4-5",
        max_tokens=4096,
        messages=[{"role": "user", "content": f"Ye document hai:\n\n{text}\n\n{get_extraction_prompt()}"}]
    )
    return message.content[0].text.strip()

def process_excel_file(file_path):
    print("Excel/CSV process ho rahi hai...")
    ext = os.path.splitext(file_path)[1].lower()
    try:
        if ext == ".csv":
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                text = f.read()
        else:
            try:
                import openpyxl
            except ImportError:
                os.system("pip install openpyxl")
                import openpyxl
            wb = openpyxl.load_workbook(file_path)
            ws = wb.active
            rows = []
            for row in ws.iter_rows(values_only=True):
                row_text = " | ".join([str(c) for c in row if c is not None])
                if row_text.strip():
                    rows.append(row_text)
            text = "\n".join(rows)
    except Exception as e:
        print(f"Read error: {e}")
        text = ""
    client = anthropic.Anthropic(api_key=ANTHROPIC_API_KEY)
    message = client.messages.create(
        model="claude-opus-4-5",
        max_tokens=4096,
        messages=[{"role": "user", "content": f"Ye data hai:\n\n{text}\n\n{get_extraction_prompt()}"}]
    )
    return message.content[0].text.strip()

def clean_json(response_text):
    if "```" in response_text:
        parts = response_text.split("```")
        for part in parts:
            if "[" in part or "{" in part:
                response_text = part
                if response_text.startswith("json"):
                    response_text = response_text[4:]
                break
    array_start = response_text.find("[")
    object_start = response_text.find("{")
    if array_start != -1 and (object_start == -1 or array_start < object_start):
        start = array_start
        end = response_text.rfind("]") + 1
    else:
        start = object_start
        end = response_text.rfind("}") + 1
    if start != -1 and end > 0:
        response_text = response_text[start:end]
    return response_text

def is_duplicate(sheet, phone, invoice_number):
    if not phone and not invoice_number:
        return False
    try:
        all_data = sheet.get_all_values()
        for row in all_data[1:]:
            if len(row) < 4:
                continue
            existing_phone = str(row[3]).strip()
            existing_invoice = str(row[7]).strip() if len(row) > 7 else ""
            if invoice_number and existing_invoice:
                if invoice_number.strip() == existing_invoice:
                    return True
            elif phone and existing_phone:
                if phone.strip() == existing_phone:
                    return True
    except:
        pass
    return False

def add_to_sheet(sheet, data, filename, file_type):
    phone = str(data.get("phone") or "").strip()
    invoice_number = str(data.get("invoice_number") or "").strip()

    if is_duplicate(sheet, phone, invoice_number):
        print(f"   ⚠️ Duplicate skip: {data.get('name', 'Unknown')} — {phone}")
        return

    # Fold new fields (gst_number, date) into extra_info so sheet columns stay unchanged
    extra_parts = []
    if data.get("gst_number"):
        extra_parts.append(f"GST: {data['gst_number']}")
    if data.get("date"):
        extra_parts.append(f"Doc Date: {data['date']}")
    if data.get("extra_info"):
        extra_parts.append(str(data["extra_info"]).strip())
    extra_info = " | ".join(extra_parts)

    now = datetime.now()
    row = [
        now.strftime("%d-%m-%Y"),
        now.strftime("%H:%M:%S"),
        str(data.get("name") or "Unknown").strip(),
        phone,
        str(data.get("email") or "").strip(),
        str(data.get("company") or "").strip(),
        str(data.get("address") or "").strip(),
        invoice_number,
        str(data.get("amount") or "").strip(),
        extra_info,
        "New Lead",
        filename,
        file_type.upper()
    ]
    sheet.append_row(row)
    print(f"   ✅ Added: {data.get('name', 'Unknown')} — {phone}")

def process_file(file_path):
    if not os.path.exists(file_path):
        print(f"❌ File nahi mili: {file_path}")
        return

    file_type = get_file_type(file_path)
    print(f"\nFile type: {file_type.upper()}")

    if file_type == "unknown":
        print("❌ Unsupported file type")
        return

    if file_type == "image":
        response = process_image_file(file_path)
    elif file_type == "pdf":
        response = process_pdf_file(file_path)
    elif file_type == "word":
        response = process_word_file(file_path)
    elif file_type == "excel":
        response = process_excel_file(file_path)

    cleaned = clean_json(response)
    data = json.loads(cleaned)
    sheet = connect_to_sheet()
    clients = data if isinstance(data, list) else [data]

    print(f"{len(clients)} entries mili hain...")
    for client in clients:
        add_to_sheet(sheet, client, os.path.basename(file_path), file_type)

    print("\n" + "="*45)
    print(f"✅ DONE! {len(clients)} processed")
    print("="*45)
    for i, c in enumerate(clients, 1):
        print(f"   {i}. {c.get('name','N/A')} — {c.get('phone','N/A')}")
    print("="*45)

if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("\n📁 Usage: python extract.py <file_path>")
    else:
        process_file(sys.argv[1])